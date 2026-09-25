from pathlib import Path
import csv, json, math, hashlib
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path('D:/Desktop/E')
WORK=BASE/'question1/paper_work'
RES=BASE/'question1/results'
OUT=BASE/'问题一_分析与数学建模_完整版.docx'
ROWS=list(csv.DictReader((RES/'sample_summary.csv').open(encoding='utf-8-sig')))
STATS=json.loads((RES/'validation.json').read_text('utf-8'))
EX=json.loads((RES/'example.json').read_text('utf-8'))
ENV=json.loads((RES/'run_environment.json').read_text('utf-8'))
FONT='C:/Windows/Fonts/simsun.ttc'
BOLD='C:/Windows/Fonts/simhei.ttf'
INK='#20252B'; BLUE='#326279'; TEAL='#397D7F'; GREY='#73818B'; LIGHT='#EEF3F5'; ORANGE='#BD7840'
def font(size,bold=False):return ImageFont.truetype(BOLD if bold else FONT,size)
def centered(draw,xy,text,size=29,fill=INK,bold=False):
    draw.text(xy,text,font=font(size,bold),fill=fill,anchor='mm',align='center',spacing=10)
def arrow(draw,a,b,color=GREY,width=4):
    draw.line([a,b],fill=color,width=width)
    dx,dy=b[0]-a[0],b[1]-a[1]; v=math.hypot(dx,dy); ux,uy=dx/v,dy/v
    draw.polygon([b,(b[0]-16*ux+7*uy,b[1]-16*uy-7*ux),(b[0]-16*ux-7*uy,b[1]-16*uy+7*ux)],fill=color)
def box(draw,rect,text,size=29):
    draw.rounded_rectangle(rect,12,fill=LIGHT,outline=GREY,width=2)
    centered(draw,((rect[0]+rect[2])/2,(rect[1]+rect[3])/2),text,size)
def figures():
    im=Image.new('RGB',(1800,680),'white');d=ImageDraw.Draw(im)
    box(d,(580,15,1220,90),'附件1原始视频与英文转写',34)
    for x,txt in [(50,'文本分词与字符偏移'),(655,'音频解码与16 kHz重采样'),(1260,'按真实时间戳采样视频')]:
        box(d,(x,155,x+490,225),txt,28)
    for x in [295,900,1505]:arrow(d,(900,95),(x,150))
    for x,txt in [(50,'DistilBERT\n词级语义768维'),(655,'openSMILE\n帧级声学25维'),(1260,'MediaPipe\n表情与头姿55维')]:
        box(d,(x,280,x+490,390),txt,29);arrow(d,(x+245,230),(x+245,274))
    box(d,(610,438,1190,516),'Wav2Vec2与CTC强制对齐\n获得每个词的起止时间',27)
    d.line([(540,190),(575,190),(575,477)],fill=GREY,width=4);arrow(d,(575,477),(605,477))
    d.line([(1145,190),(1215,190),(1215,477)],fill=GREY,width=4);arrow(d,(1215,477),(1195,477))
    d.line([(900,398),(900,416),(1240,416)],fill=GREY,width=4);arrow(d,(1240,416),(1240,554))
    for x in [295,1505]:arrow(d,(x,398),(x,554))
    arrow(d,(900,520),(900,554))
    box(d,(50,562,1750,650),'按词区间重叠加权聚合  →  文本 L×768  语音 L×50  视觉 L×110\n保存时间映射  有效掩码  置信度  源帧与处理日志',29)
    im.save(WORK/'flow.png')

    im=Image.new('RGB',(1800,390),'white');d=ImageDraw.Draw(im)
    centered(d,(900,35),'以重复字符 ll 为例的CTC状态转移示意',32,bold=True)
    xs=[160,515,870,1225,1580]; labs=['blank','l (1)','blank','l (2)','blank']
    for x,t in zip(xs,labs):
        d.ellipse((x-78,105,x+78,261),fill=LIGHT,outline=BLUE,width=3);centered(d,(x,182),t,32)
        d.arc((x-34,67,x+34,136),180,530,fill=GREY,width=3)
    for a,b in zip(xs[:-1],xs[1:]):arrow(d,(a+83,182),(b-85,182),BLUE)
    d.line([(515,270),(515,309),(1225,309),(1225,270)],fill=ORANGE,width=3)
    centered(d,(870,340),'相同字符之间禁止跨越blank直接跳转',28,ORANGE)
    im.save(WORK/'ctc.png')

    im=Image.new('RGB',(1800,380),'white');d=ImageDraw.Draw(im)
    x0,x1=300,1620
    def xx(v):return x0+(x1-x0)*v/3
    centered(d,(900,28),'时间重叠加权示例  I = [0.5, 2.0] 秒',31,bold=True)
    for y,label in [(112,'目标词区间'),(217,'源帧支持区间')]:d.text((35,y),label,font=font(28),fill=INK)
    d.rectangle((xx(.5),95,xx(2),150),fill='#D5E7EB',outline=BLUE,width=3)
    for a,b,t,c in [(0,1,'x(1) = 2  有效',TEAL),(1,2,'x(2) = 4  有效',BLUE),(2,3,'x(3) = 99  无效',GREY)]:
        d.rectangle((xx(a)+3,196,xx(b)-3,253),outline=c,width=3,fill=LIGHT);centered(d,((xx(a)+xx(b))/2,224),t,26)
    d.line((x0,283,x1,283),fill=INK,width=2)
    for t in [0,.5,1,2,3]:
        d.line((xx(t),278,xx(t),293),fill=INK,width=2);centered(d,(xx(t),319),str(t),25)
    centered(d,(900,361),'权重为 0.5、1.0、0    加权均值为 10/3    加权标准差约 0.943',27)
    im.save(WORK/'pool.png')

    im=Image.new('RGB',(1800,1060),'white');d=ImageDraw.Draw(im)
    def panel(rect,title,vals,bins):
        l,t,r,b=rect; centered(d,((l+r)/2,t+22),title,30,bold=True)
        px=l+80;py=b-70;pw=r-px-25;ph=b-t-150
        count,edges=np.histogram(vals,bins=bins); ymax=max(count)+2
        for k in range(0,int(ymax)+1,5):
            y=py-ph*k/ymax;d.line((px,y,px+pw,y),fill='#DFE4E8',width=1);d.text((px-12,y),str(k),anchor='rm',font=font(22),fill=GREY)
        for i,c in enumerate(count):
            a=px+pw*i/len(count);bb=px+pw*(i+1)/len(count);d.rectangle((a+4,py-ph*c/ymax,bb-4,py),fill=BLUE)
            centered(d,((a+bb)/2,py+25),f'{edges[i]:g}–{edges[i+1]:g}',20)
        d.line((px,py,px+pw,py),fill=INK,width=2)
    panel((10,0,890,480),'（a）解码时长分布  单位为秒',[float(x['duration_s']) for x in ROWS],[0,5,10,15,20,25,30])
    panel((910,0,1790,480),'（b）词数分布  单位为词',[int(x['word_count']) for x in ROWS],[0,10,20,30,40,50,60,70])
    centered(d,(450,520),'（c）词级有效覆盖率',30,bold=True)
    for i,(lab,v) in enumerate([('时间区间',1),('语音',1),('视觉',1601/1931)]):
        y=600+i*95;d.text((35,y+22),lab,font=font(28),fill=INK,anchor='lm');d.rectangle((225,y,775,y+46),fill=LIGHT);d.rectangle((225,y,225+550*v,y+46),fill=TEAL);d.text((790,y+22),f'{v:.1%}',font=font(25),fill=INK,anchor='lm')
    centered(d,(1340,520),'（d）样本平均对齐分数与人脸检出率',29,bold=True)
    px=1000;py=965;pw=690;ph=350
    for k in range(6):
        v=k/5;y=py-v*ph;x=px+v*pw
        d.line((px,y,px+pw,y),fill='#DFE4E8',width=1);d.text((px-13,y),f'{v:.1f}',font=font(21),fill=GREY,anchor='rm');centered(d,(x,py+24),f'{v:.1f}',21)
    for row in ROWS:
        x=px+float(row['alignment_confidence_mean'])*pw;y=py-float(row['face_detection_rate'])*ph;c=ORANGE if float(row['alignment_confidence_mean'])<.1 else BLUE
        d.ellipse((x-6,y-6,x+6,y+6),fill=c)
    d.line((px+.1*pw,py-ph,px+.1*pw,py),fill=ORANGE,width=2)
    centered(d,(1340,1028),'横轴为对齐分数  纵轴为人脸检出率',24)
    centered(d,(450,965),'覆盖率衡量可用性  不等于对齐准确率',25,GREY)
    im.save(WORK/'quality_cn.png')

    source=Image.open(RES/'example_frames.jpg')
    # Reflow the six existing source-video panels to improve print readability.
    im=Image.new('RGB',(1500,660),'white')
    sw=source.width/6
    for i in range(6):
        crop=source.crop((round(i*sw),130,round((i+1)*sw),337))
        crop=crop.resize((480,265),Image.Resampling.LANCZOS)
        xx0=(i%3)*500+(500-crop.width)//2;yy=(i//3)*330
        im.paste(crop,(xx0,yy))
        lab=['I  0.10 s  帧3','think  0.30 s  帧9','justice  1.60 s  帧48','extremely  2.20 s  帧66','for  2.90 s  帧87','people  3.50 s  帧105'][i]
        centered(ImageDraw.Draw(im),((i%3)*500+250,yy+302),lab,25)
    im.save(WORK/'frames_grid.png')

def setfont(run,size=11,bold=False,east='宋体'):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    pr=run._element.get_or_add_rPr();rf=pr.rFonts
    if rf is None:rf=OxmlElement('w:rFonts');pr.insert(0,rf)
    rf.set(qn('w:eastAsia'),east)

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(2.1);sec.left_margin=sec.right_margin=Cm(2.4)
sec.header_distance=Cm(.7);sec.footer_distance=Cm(.9)
for name in ['Normal','Body Text','Title','Heading 1','Heading 2','Heading 3','Caption']:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
    st._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    st.paragraph_format.line_spacing=Pt(17)
    st.paragraph_format.space_after=Pt(4)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.first_line_indent=Pt(22)
for name,size in [('Title',17),('Heading 1',14),('Heading 2',11.5)]:
    st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
    st._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体')
    st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.keep_with_next=True
    st.paragraph_format.space_before=Pt(7);st.paragraph_format.space_after=Pt(6)
doc.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for st in doc.styles:
    for border in list(st._element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;foot.paragraph_format.first_line_indent=Pt(0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
doc.core_properties.title='问题一 多模态情感特征提取与时序对齐的分析和数学建模'
doc.core_properties.author='';doc.core_properties.last_modified_by=''
doc.core_properties.subject='附件1全量样本的特征提取 时序对齐 质量核验'

def p(text,size=11,indent=True):
    q=doc.add_paragraph();q.paragraph_format.first_line_indent=Pt(22 if indent else 0)
    q.paragraph_format.widow_control=True
    q.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    r=q.add_run(text);setfont(r,size);return q
def h(text,level=1):return doc.add_paragraph(text,'Heading '+str(level))
def page():doc.add_page_break()
def cap(text):
    q=doc.add_paragraph();q.paragraph_format.first_line_indent=Pt(0);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_before=Pt(3);q.paragraph_format.space_after=Pt(5);q.paragraph_format.line_spacing=Pt(13)
    setfont(q.add_run(text),9.5);return q
def pic(path,width=15.7):
    q=doc.add_paragraph();q.paragraph_format.first_line_indent=Pt(0);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.line_spacing=1;q.paragraph_format.space_after=Pt(0);q.paragraph_format.keep_with_next=True
    q.add_run().add_picture(str(path),width=Cm(width))
def mr(text,plain=False):
    r=OxmlElement('m:r')
    if plain:
        pr=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');pr.append(sty);r.append(pr)
    wp=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'22');wp.append(sz);r.append(wp)
    t=OxmlElement('m:t');t.text=str(text);r.append(t);return r
def contents(parent,items):
    if not isinstance(items,(list,tuple)):items=[items]
    for item in items:parent.append(mr(item) if isinstance(item,str) else item)
    return parent
def sub(base,index):
    el=OxmlElement('m:sSub');contents(el.append if False else el,[])
    contents(el_make(el,'m:e'),base);contents(el_make(el,'m:sub'),index);return el
def el_make(parent,tag):
    x=OxmlElement(tag);parent.append(x);return x
def sup(base,power):
    el=OxmlElement('m:sSup');contents(el_make(el,'m:e'),base);contents(el_make(el,'m:sup'),power);return el
def ss(base,index,power):
    el=OxmlElement('m:sSubSup');contents(el_make(el,'m:e'),base);contents(el_make(el,'m:sub'),index);contents(el_make(el,'m:sup'),power);return el
def frac(a,b):
    el=OxmlElement('m:f');contents(el_make(el,'m:num'),a);contents(el_make(el,'m:den'),b);return el
def summ(index,body):
    el=OxmlElement('m:nary');pr=el_make(el,'m:naryPr')
    for tag,value in [('chr','∑'),('limLoc','subSup'),('supHide','1')]:
        v=el_make(pr,'m:'+tag);v.set(qn('m:val'),value)
    contents(el_make(el,'m:sub'),index);el_make(el,'m:sup');contents(el_make(el,'m:e'),body);return el
def sqrt(body):
    el=OxmlElement('m:rad');pr=el_make(el,'m:radPr');v=el_make(pr,'m:degHide');v.set(qn('m:val'),'1');el_make(el,'m:deg');contents(el_make(el,'m:e'),body);return el
def lim(name,condition):
    el=OxmlElement('m:limLow');contents(el_make(el,'m:e'),mr(name,True));contents(el_make(el,'m:lim'),condition);return el
def eq(text,n):
    q=doc.add_paragraph();q.paragraph_format.first_line_indent=Pt(0);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.line_spacing=1.05;q.paragraph_format.space_after=Pt(6);q.paragraph_format.space_before=Pt(3)
    expressions={
      1:['t = ',mr('PTS',True),' × ',mr('time_base',True),' − ',sub('t','0')],
      2:[ss('x','i','t'),' = ',frac('1',['|',sub('S','i'),'|']),summ(['j ∈ ',sub('S','i')],sub('h','j')), '，  ',sup('X','t'),' ∈ ',sup('ℝ','L×768')],
      3:[ss('m','j','a'),' = 1{',sub('ρ','j'),' ≥ 0.95} · 1{',ss('x','j','a'),mr('各维均为有限值',True),'}'],
      4:['R = U ',mr('diag',True),'(1, 1, ',mr('det',True),'(U',sup('V','T'),')) ',sup('V','T')],
      5:[ss('x','j','v'),' = [',sub('b','j,1'),', …, ',sub('b','j,52'),', ',sub('θ','j,x'),', ',sub('θ','j,y'),', ',sub('θ','j,z'),'] ∈ ',sup('ℝ','55')],
      6:[sup('π','*'),' = ',lim('arg max','π ∈ 𝒜(ℓ)'),summ('t',[mr('log',True),' ',sub('p','t'),'(',sub('π','t'),')'])],
      7:[sub('δ','t'),'(s) = ',mr('log',True),' ',sub('p','t'),'(',sub('z','s'),') + ',lim('max','u ∈ P(s)'),sub('δ','t−1'),'(u)'],
      8:[sub('b','k'),' = ',mr('clip',True),'(',sub('c','0'),' + (k − 1/2)Δt, 0, D)'],
      9:[sub('c','k'),' = ',mr('exp',True),'[',frac('1',['|',sub('F','k'),'|']),summ(['t ∈ ',sub('F','k')],[mr('log',True),' ',sub('p','t'),'(',sub('ℓ','k'),')']),']'],
      10:[sub('q','i'),' = ',mr('exp',True),'[',frac('1',['|',sub('C','i'),'|']),summ(['k ∈ ',sub('C','i')],[mr('log',True),'(',mr('max',True),'(',sub('c','k'),', ',sup('10','−30'),'))']),']'],
      11:[sub('w','ij'),' = ',mr('max',True),'(0, ',mr('min',True),'(',sub('e','i'),',',sub('f','j'),') − ',mr('max',True),'(',sub('s','i'),',',sub('g','j'),')) · ',sub('m','j')],
      12:[sub('W','i'),' = ',summ('j',sub('w','ij')),'，  ',sub('μ','i'),' = ',frac('1',sub('W','i')),summ('j',[sub('w','ij'),sub('x','j')])],
      13:[sub('σ','i'),' = ',sqrt([frac('1',sub('W','i')),summ('j',[sub('w','ij'),sup(['(',sub('x','j'),' − ',sub('μ','i'),')'],'2')])])]
    }
    mathp=OxmlElement('m:oMathPara');m=OxmlElement('m:oMath');contents(m,expressions[n]);m.append(mr('    ('+str(n)+')',True));mathp.append(m);q._p.append(mathp)
    return q
def table(headers,rows,widths=None,size=9.5):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths is None:widths=[16.2/len(headers)]*len(headers)
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    marg=OxmlElement('w:tblCellMar')
    for side,val in [('top',45),('bottom',45),('left',65),('right',65)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');marg.append(e)
    pr.append(marg)
    for j,txt in enumerate(headers):t.rows[0].cells[j].text=txt
    for row in rows:
        cells=t.add_row().cells
        for j,txt in enumerate(row):cells[j].text=str(txt)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
        if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,cell in enumerate(row.cells):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEFF1');cell._tc.get_or_add_tcPr().append(sh)
            for q in cell.paragraphs:
                q.paragraph_format.first_line_indent=Pt(0);q.paragraph_format.space_after=Pt(0);q.paragraph_format.space_before=Pt(0)
                q.paragraph_format.line_spacing=Pt(size+3)
                q.alignment=WD_ALIGN_PARAGRAPH.CENTER if j==0 or len(str(cell.text))<12 else WD_ALIGN_PARAGRAPH.LEFT
                for r in q.runs:setfont(r,size,i==0)
    return t

figures()
# Page 1
doc.add_paragraph('多模态情感特征提取与时序对齐',style='Title')
q=doc.add_paragraph('问题一的分析与数学建模');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.first_line_indent=Pt(0);setfont(q.runs[0],12)
p('针对附件1中100条英文视频，建立以词级时间区间为统一索引的三模态特征提取模型。模型将文本上下文表示、声学描述符、面部表情与头部姿态映射到同一词序列，同时保留原始帧级特征、时间映射和质量掩码。现有运行结果覆盖全部样本，共生成1931个词级位置，语音有效覆盖率为100%，视觉有效覆盖率为82.91%；结构与聚合复算检查未发现错误。上述指标反映数据生成与可用性，不代表人工验证的对齐准确率。')
p('关键词  多模态特征提取；CTC强制对齐；Viterbi算法；时间重叠加权；可追溯时序表示',10,False)
h('1 问题分析')
h('1.1 任务目标与研究边界',2)
p('问题一要求从原始素材出发，给出特征定义、提取方法以及跨模态时序组织规则。与直接使用预计算特征不同，本问需要说明每个数值向量来自哪段文本、哪段语音及哪些视频帧。因此，模型的主要输出是可核验的时序特征文件，而不是情感极性或情感强度的预测结果。本文仅使用赛题提供的样本，预训练工具用于冻结参数的基础处理，不使用情感标签训练或微调特征提取器。')
h('1.2 异构模态带来的建模困难',2)
p('文本以离散词语承载语义，音频是连续采样信号，视频则以帧记录面部状态。三者既有采样频率差异，也有表示维度和有效观测范围差异。若仅将序列裁剪到相同长度，则相同下标未必对应同一段内容；若对整句直接求平均，则会丢失局部情感变化和证据位置。另一方面，画外音、遮挡及转写不匹配可能使某些模态的局部信息不可用，需将“没有观测”与“观测值恰为零”明确区分。')
h('1.3 以词为时序锚点的解决思路',2)
p('本文先独立提取各模态特征，再利用已知转写与音频进行强制对齐，为每个词确定时间区间，最后按区间重叠长度聚合语音和视觉特征。词级索引兼顾语言可读性与时间可定位性，便于将数值输出回查至原始素材。对齐低分、无人脸及多人脸情况通过独立标记保留，不通过删除样本提高表面覆盖率。')
cap('表1  问题一的要求与对应实现')
table(['题目要求','实现方式','核验材料'],[
('样本覆盖完整','100条样本逐一匹配并保留','清单、特征文件、源视频哈希'),
('时序组织可核验','词区间、源帧索引、有效长度与掩码','逐词明细与典型样本图'),
('方法可复现','固定配置、模型修订号与处理日志','环境快照、代码与运行说明')],[3.1,6.1,7],9.5)

# Page 2
page();h('2 模型假设与符号说明')
h('2.1 建模假设及适用条件',2)
p('第一，容器中的音频、视频时间戳可作为建立公共时间轴的依据。程序使用实际解码时间戳而非单纯按帧号除以帧率换算；存在音频缺口时保留观测掩码。该假设不排除源视频本身已有音画偏移，因此所得对应关系仍受素材质量限制。')
p('第二，给定转写可作为强制对齐的目标序列，但不假定其与音频始终一致。数字读法采用显式映射，不修改原始文本；无法建立合法路径时标记失败，低分路径保留复核信息。该处理把文本一致性作为需要检查的条件，避免把强制匹配误写为已证实的正确匹配。')
p('第三，源帧特征在其短时间支持区间内可由当前帧代表，因而能用时间重叠长度进行加权。该近似适合构造词级统计量，但会平滑词内变化；原始帧级数据同时保存，允许后续使用更细粒度方法重新分析。')
p('第四，视觉端每帧选择最大人脸作为候选表达主体。这是确定性选择规则，而不是已验证的说话人识别假设。多人脸与低检出率场景需回看，不能仅凭人脸面积确定该人正在发声。')
h('2.2 主要符号',2)
cap('表2  数学符号及其含义')
table(['符号','定义或含义'],[
('n，N','样本索引与样本总数，本问N = 100'),
('Dₙ，Lₙ','第n条样本的解码时长与词数'),
('wᵢ，Sᵢ','第i个词及其对应的文本子词下标集合'),
('Iᵢ = [sᵢ,eᵢ)','第i个词在公共时间轴上的区间'),
('Jⱼ = [gⱼ,fⱼ)','第j个声学帧或视觉帧的时间支持区间'),
('xⱼ，mⱼ','源帧特征向量与有效观测指示量'),
('pₜ(c)，ℓ','第t个声学时间步的字符概率与目标字符序列'),
('zₛ，δₜ(s)','CTC增广状态的符号及到达状态的最优对数分数'),
('qᵢ，rᵢ','词级对齐分数及低分复核标记'),
('wᵢⱼ，Wᵢ','词与源帧的重叠权重及该词的总权重'),
('μᵢ，σᵢ','重叠加权均值向量与逐维标准差向量'),
('Xᵗ，Xᵃ，Xᵛ','词级文本、语音、视觉特征矩阵')],[4.2,12],10)
p('为简化表达，后文在单条样本内省略样本下标n。时间单位统一为秒，区间采用左闭右开约定。模态掩码为True表示存在有效观测贡献；补齐掩码为True表示批量输入中人为增加的位置，二者含义相反。',10)

# Page 3
page();h('3 总体框架与原始数据处理')
pic(WORK/'flow.png');cap('图1  从原始素材到词级三模态特征的处理流程')
h('3.1 样本编号与输入完整性',2)
p('程序读取附件1的label-100.xlsx，以video_id与clip_id共同定位视频，并构造唯一标识video_id$_$clip_id。读取时核验样本总数为100、编号无重复、表格视频集合与磁盘视频集合一致，并检查标签取值与正负中性含义。原转写和标签原样写入元数据；标签只用于一致性核验及结果登记，不进入特征提取或对齐的目标函数。')
h('3.2 公共时间轴与观测掩码',2)
p('使用PyAV解码，取容器起始时刻为时间原点。若某帧具有呈现时间戳PTS及对应时间基，则其相对时刻按式（1）计算。音频通过重采样器转换为16 kHz单声道，并根据重采样帧的时间戳放入公共采样数组，而不是忽略偏移后直接拼接。')
eq('t = PTS × time_base − t₀',1)
p('音频数组初始化为零，同时建立同长度的观测布尔数组。写入真实解码数据的位置记为已观测，其余位置仍为未观测。零值在此只承担存储占位作用，后续声学有效性判断还需检查观测比例。因此，容器缺口与实际静音被保留为不同状态。缺失声学上下文仍可能影响强制对齐，这属于后续复核范围。')
h('3.3 文本清理与时间口径',2)
p('分词规则保留英文缩写中的撇号、带千分位的数字及序数词，并登记原文字符偏移。标点保留在原始文本中，但不作为独立词级特征位置；连字符连接的词按规则拆分。音频对齐使用单独的字符规范化副本，避免对赛题原文作不可追溯修改。')
p('现有处理结果的解码总时长为787.5022秒，单条范围为2.2570至29.2881秒。该范围与题面所述范围不同，本文不据此替换或裁剪视频，而在全量表中统一报告程序记录的解码时间轴长度。时长定义和源文件哈希共同保留，便于对差异进一步复核。')

# Page 4
page();h('4 三模态情感相关特征的构建')
h('4.1 文本上下文表示',2)
p('文本采用冻结的DistilBERT英文编码器[1]。将分词结果作为预分词输入，模型进一步产生子词序列和768维最后隐层输出。一个原词可能对应多个子词，若直接使用子词位置将与音频词区间不一致，因此依据分词器返回的word_ids建立映射，对属于同一词的隐层向量取均值：')
eq('xᵢᵗ = (1 / |Sᵢ|) ∑ⱼ∈Sᵢ hⱼ，   Xᵗ ∈ ℝᴸˣ⁷⁶⁸',2)
p('程序按最多100个原词分块编码，不静默截断；若子词数量超过模型允许长度，则报错要求减小块大小。当前样本最长65词，均在单个词块内处理。该表示保留块内上下文，但预处理时去掉了作为独立标点的输入，不能声称完全保留原句全部语言线索。语义向量也不直接对应正负情绪概率，其情感判别价值需在下游任务中检验。')
h('4.2 语音低层声学描述符',2)
p('语音使用openSMILE 2.6.0的eGeMAPSv02低层描述符配置[3]，每个声学帧输出25维向量。选用帧级描述符是为了保留后续按词聚合的时间自由度；该输出并非整句88维统计特征。特征起止时刻直接采用工具返回的帧区间，避免另行假定固定窗口后造成错位。')
cap('表3  25维声学描述符的组成')
table(['特征组','主要内容','维数'],[
('能量与频谱','响度、alphaRatio、Hammarberg指数、两段谱斜率、谱通量','6'),
('倒谱','第1至第4阶MFCC','4'),
('基频与音质','半音基频、局部抖动、微颤、谐噪比HNR','4'),
('谐波差异','H1–H2、H1–A3相关对数相对量','2'),
('共振峰','前三个共振峰的频率、带宽及相对幅度','9'),
('合计','按feature_names.json固定列顺序保存','25')],[3,11.5,1.7],9.5)
p('为避免未观测音频对统计量造成污染，代码计算每帧区间内真实音频采样点的占比ρⱼ，同时检查各维数值是否有限。仅当观测比例不低于95%且数值全部有限时，该声学帧参与词级聚合：')
eq('mⱼᵃ = 1{ρⱼ ≥ 0.95} · 1{xⱼᵃ各维均为有限值}',3)
p('这项规则只排除缺失或数值异常帧，并不将所有背景噪声或非语音片段识别为无效。基频为零也不能在此处简单等同于缺失。声学特征的物理含义与质量掩码应结合解释，最终通过词区间内均值和标准差得到50维词级表示。')

# Page 5
page();h('4 三模态情感相关特征的构建')
h('4.3 面部表情与头部姿态',2)
p('视觉端按10 Hz目标频率选取具有真实PTS的视频帧。程序依次解码，当帧时刻到达下一个目标采样位置时保留该帧，并据实际时刻更新下一采样门限；因此采样间隔可随原视频时间戳变化。图像宽度超过640像素时等比例缩小，限制的是宽度而非最长边。')
p('采用MediaPipe Face Landmarker[4]检测至多3张人脸，检测与存在置信度阈值均为0.5。多张人脸时依据关键点包围框面积选择最大者。对被选人脸读取52维blendshape系数与面部变换矩阵；前者表示表情混合形状系数，不等同于FACS动作单元，也不能直接视为情感类别。')
pic(WORK/'frames_grid.png',14.8);cap('图2  典型样本的原始视频采样帧及可回查位置')
p('设变换矩阵左上角3×3部分为A。为从可能含尺度因素的矩阵中得到合法旋转，先作奇异值分解A = UΣVᵀ，再校正行列式符号，得到旋转矩阵R，并转为xyz顺序的欧拉角。角度使用弧度制：')
eq('R = U diag(1, 1, det(UVᵀ)) Vᵀ',4)
eq('xⱼᵛ = [bⱼ,₁, …, bⱼ,₅₂, θⱼ,ₓ, θⱼ,ᵧ, θⱼ,𝓏] ∈ ℝ⁵⁵',5)
p('相邻采样时刻的中点构成帧支持区间边界，且每帧向左右的支持范围均限制在0.075秒以内，首末区间再限制于视频时长内。该规则防止跨越较大解码缺口。检测失败或数值无效时置零并记mⱼᵛ = 0；同时保留人脸数、选中框和原始帧号，使视觉选择及缺失可回查。图2中的采样时刻是帧PTS，是否参与某词聚合由支持区间重叠决定，不能只比较中心点是否落在词区间内。')

# Page 6
page();h('5 基于CTC的词级强制对齐模型')
h('5.1 目标序列与最大概率路径',2)
p('使用Wav2Vec2英文CTC声学模型[2]获得音频各时间步的字符对数概率。转写先大写化，对登记过的数字与序数词进行读法展开，例如10th展开为TENTH；词间插入词分隔符“|”。需要区分词分隔符与CTC blank：前者是目标字符，后者用于状态停留和重复字符分离。原文与字符所属原词的映射同时保留。')
eq('π* = arg max π∈𝒜(ℓ) ∑ₜ log pₜ(πₜ)',6)
p('式（6）中𝒜(ℓ)表示能够折叠为目标序列ℓ的合法CTC路径集合。此处求解的是固定声学后验下的最优受约束路径，不涉及重新训练声学模型。将目标扩展为z = (blank, ℓ₁, blank, …, ℓₖ, blank)，用δₜ(s)记录到达第s个状态的最大对数分数。')
h('5.2 Viterbi动态规划与回溯',2)
eq('δₜ(s) = log pₜ(zₛ) + max u∈P(s) δₜ₋₁(u)',7)
p('前驱集合P(s)始终包含合法范围内的s与s−1；只有zₛ不是blank且zₛ与zₛ₋₂不同，才允许加入s−2。初始化时仅首个blank与首个目标字符可达，其余状态设为负无穷。终止时从最后一个字符或末尾blank中选择较大分数，依据记录的转移回溯完整路径。若终止状态不可达，则报告对齐失败。')
pic(WORK/'ctc.png',15.1);cap('图3  重复字符必须经过blank的完整CTC转移约束')
h('5.3 声学帧到秒的换算',2)
p('Wav2Vec2卷积前端的总步幅为320个采样点、感受野为400个采样点。对16 kHz输入，时间步长Δt = 0.020秒，首个感受野中心c₀ = 399/(2×16000)。以中心之间的中点构造帧单元边界并裁剪至[0,D]：')
eq('bₖ = clip(c₀ + (k − 1/2)Δt, 0, D)',8)
p('若某词首字符起于帧u，末字符止于帧v（含v），则词区间为[bᵤ,bᵥ₊₁)。这样得到的20毫秒步长是离散时间分辨率，不能据此声称词边界误差小于20毫秒。重复字符拓扑保证路径合法性，但无法消除转写错误或声学识别偏差。')

# Page 7
page();h('6 词级聚合与质量标记模型')
h('6.1 对齐置信度的定义',2)
p('对目标字符k，设Fₖ为回溯路径中发射该字符的声学帧集合，先对字符发射概率取几何平均得到cₖ；再对属于词i的字符集合Cᵢ取几何平均得到qᵢ。该两级平均与代码一致，使每个目标字符在词分数中具有相同权重，而非让发射帧较多的字符获得更高权重。')
eq('cₖ = exp[(1/|Fₖ|)∑ₜ∈Fₖ log pₜ(ℓₖ)]',9)
eq('qᵢ = exp[(1/|Cᵢ|)∑ₖ∈Cᵢ log(max(cₖ,10⁻³⁰))]',10)
p('定义复核标记rᵢ = 1{qᵢ < 0.1 或对齐失败}。阈值0.1是当前配置中的检查门限，未经人工边界真值校准，不是错误概率阈值。低分词仍保留时间区间与特征；没有合法路径时保留文本位置，词区间与对齐掩码置为无效，并保存原始声学、视觉数据。')
h('6.2 按时间重叠进行加权统计',2)
eq('wᵢⱼ = max(0, min(eᵢ,fⱼ) − max(sᵢ,gⱼ)) · mⱼ',11)
eq('Wᵢ = ∑ⱼwᵢⱼ，  μᵢ = (1/Wᵢ)∑ⱼwᵢⱼxⱼ',12)
eq('σᵢ = √[(1/Wᵢ)∑ⱼwᵢⱼ(xⱼ − μᵢ)²]',13)
p('式（12）与式（13）仅在Wᵢ > 0时计算，平方和开方按特征维逐项进行。均值描述词内平均状态，标准差补充词内变化幅度。将二者拼接后，语音由25维变为50维，视觉由55维变为110维；文本仍保持768维。累加过程使用float64，保存时转为float32，以减少大数值常量描述符的舍入误差。')
pic(WORK/'pool.png',15.4);cap('图4  时间重叠加权与无效帧排除的数值示例')
h('6.3 缺失与未分配区间',2)
p('当Wᵢ = 0时，输出零向量并将该词的模态掩码置False；Wᵢ > 0只说明存在有效贡献，不保证覆盖整个词区间。词间未分配区间取所有有效词区间在[0,D]内的补集，单独保存而不自动称为静音。帧级聚合会压缩词内先后顺序，因此原始帧特征及其时间戳一并保留，以支持复算与更细的时序分析。')

# Page 8
page();h('7 算法实现与特征文件规范')
h('7.1 单条样本的执行逻辑',2)
p('主程序由process_one串联处理。首先按真实时间戳解码音频并分词；随后获得词级文本向量与强制对齐区间；再提取帧级声学和视觉特征，调用统一的区间聚合函数；最后计算质量提示，检查输出数值有限性，并保存压缩数组和元数据。异常采用逐样本日志记录，全局失败计数使运行结束状态可检查。')
cap('表4  核心函数与数学模型的对应关系')
table(['代码位置','功能','对应对象'],[
('read_manifest','读取并核验输入清单','样本与视频的一一对应'),
('media_audio','解码、重采样及观测登记','式（1）与公共时间轴'),
('text_features','子词隐层向量平均','式（2）'),
('audio_features / visual_features','帧级提取与有效性检查','式（3）至式（5）'),
('align / ctc_viterbi','声学后验、动态规划与回溯','式（6）至式（10）'),
('pool_intervals','时间重叠加权均值与标准差','式（11）至式（13）'),
('collate','按批内最长序列补齐','补齐掩码与有效长度')],[5.4,6.1,4.7],9)
h('7.2 文件字段与追溯关系',2)
p('每条样本输出一个NPZ数值文件和一个JSON元数据文件。NPZ保留三模态词级矩阵、原始帧特征及掩码；JSON保存原始文本、标签、字符与子词映射、每词参与聚合的源帧索引、源文件和输出文件哈希、处理指纹与告警。由“词索引—时间区间—源帧索引—视频路径”可定位原始证据。')
cap('表5  主要输出字段与形状')
table(['字段','形状或类型','含义'],[
('text / audio / vision','L×768 / L×50 / L×110','三模态词级特征'),
('word_intervals','L×2','每词起止时间'),
('word_char_offsets','L×2','原文字符起止偏移'),
('alignment_confidence','L','对齐分数，未校准'),
('各模态mask与alignment_mask','L个布尔值','观测有效性及路径有效性'),
('alignment_review_mask','L个布尔值','低分或失败复核标记'),
('audio_raw / vision_raw','Fₐ×25 / Fᵥ×55','帧级特征与配套时间字段'),
('valid_length / padding_mask','整数 / L个布尔值','实际词数与补齐标记')],[6.3,4.5,5.4],9)
h('7.3 变长存储和计算复杂度',2)
p('磁盘保存完整变长序列，不强制截断为50词。批量读取时按该批最长词数补零，补齐掩码设为True，模态有效掩码设为False。若CTC时间步数为T、目标字符数为K，动态规划的时间复杂度为O(TK)，回溯表空间为O(TK)；当前区间聚合逐词检查源帧，复杂度为O(L(Fₐ+Fᵥ))，适用于本问短视频规模。')

# Page 9
page();h('8 全量结果与质量分布')
p('对附件1的100条样本完成提取，共保留1931个词。每条样本均有词级时间区间和有效语音贡献，视觉有效词为1601个，另有330个词未获得有效人脸特征。图5展示解码时长、词数和质量指标的实测分布。')
pic(WORK/'quality_cn.png');cap('图5  全量100条样本的规模与特征可用性分布')
cap('表6  全量结果的主要统计量')
table(['指标','结果','统计口径'],[
('样本完成数','100 / 100','成功生成特征与元数据'),
('总时长与总词数','787.5022秒；1931词','解码时间轴；完整分词序列'),
('词级时间区间与语音覆盖','100%；100%','均为1931 / 1931'),
('视觉有效词比例','82.91%','1601 / 1931'),
('采样帧有效人脸比例','81.91%','6462 / 7889'),
('低分词与低均分样本','342词；18条样本','词分数或样本均分低于0.1'),
('含复核提示样本','69 / 100','各类告警的样本并集'),
('压缩特征体积','14.28 MiB','仅100个NPZ文件')],[5.4,4.5,6.3],9.5)
p('样本词数从5至65不等，其中1条超过50词，说明完整变长存储能够避免固定长度截断。低分词占17.71%，需要结合回听判断原因；另外20个采样帧存在多人脸。不同告警可以同时出现在同一条样本中，不能直接相加解释为告警样本数。')
p('视觉未覆盖可能与无人脸、遮挡或检测失败等因素有关，但现有数据仅支持“缺少有效视觉贡献”的结论，尚未逐条人工归因。图5（d）展示两种质量信号的联合分布，不构成因果关系或性能提升证明。完整逐样本数值见表9。')

# Page 10
page();h('9 典型样本的时序对应验证')
p('按原表顺序，选择首条全部词可对齐、无低分词、人脸有效率大于80%且无多人脸采样帧的样本−THoVjtIkeU$_$2。其解码时长约4.29秒，转写为“I think social justice is extremely important for young people.”。选择规则只用于展示，不能代替全量质量评价。')
pic(RES/'example_alignment.png',15.7);cap('图6  典型样本的词区间 音频波形 基频描述符与面部系数')
p('图6横轴为视频起点后的秒数。第一行是强制对齐生成的词区间，后续行分别展示音频波形、半音基频描述符及部分面部系数。词区间之间的空白仅表示未分配给目标词的时段。图2进一步提供可回看的采样帧，与表7共同构成从文本到源素材的对应证据。',10)
cap('表7  典型样本的逐词对齐结果')
table(['词','起点/s','终点/s','分数','语音帧数','视觉帧数'],[
    (x['word'],f"{x['start_s']:.3f}",f"{x['end_s']:.3f}",f"{x['score']:.3f}",x['audio_frames'],x['visual_frames']) for x in EX['word_rows']
],[3.7,2.4,2.4,2.3,2.7,2.7],9)
p('例如important对应2.422至2.782秒，由38个有效声学帧和5个有效视觉支持区间贡献，最终生成768/50/110维三模态向量。people分数约0.342，虽然低于多数其他词，但高于0.1门限，不属于当前规则下的低分词；仅凭该分数不能推断发音截断。上述对应关系属于可视核验展示，尚无人工逐词边界真值用于计算时间误差。',10)

# Page 11
page();h('10 结果核验与可复现性')
h('10.1 结构核验和独立数值复算',2)
p('汇总程序对全部特征检查样本身份、原文与标签一致性、输出文件哈希、有效长度、特征形状、有限数值、字符偏移回读、词区间正长度与单调性、视频PTS顺序。随后从保存的帧级特征重新计算词级聚合，核对结果、掩码和源帧索引。数值比较采用相对容差与绝对容差均为2×10⁻⁵的allclose规则，并非声称所有绝对误差都小于2×10⁻⁵。')
p('现有validation.json记录结构与数值检查错误列表为空。另有6项算法检查覆盖重复字符blank约束、不可行路径、无效帧排除与加权聚合、字符偏移与区间补集、数字分词以及超过50词的批量补齐。结构正确和复算一致只验证实现与保存规则，不验证词语的真实发音边界或说话人身份。')
h('10.2 固定配置与环境记录',2)
cap('表8  运行环境与主要参数')
table(['组件或配置','当前记录','用途'],[
('Python / NumPy / SciPy','3.12.14 / 1.26.4 / 1.13.1','数值计算与旋转处理'),
('PyTorch / Transformers','2.8.0 / 4.57.6','冻结模型推理'),
('PyAV / openSMILE','14.2.0 / 2.6.0','音视频解码与声学提取'),
('MediaPipe','0.10.21','人脸与表情系数提取'),
('DistilBERT修订号','12040accade4…','完整修订号见环境快照'),
('Wav2Vec2修订号','22aad52d435e…','完整修订号见环境快照'),
('音频采样率 / 视频采样率','16000 Hz / 10 Hz','公共时间轴与采样配置'),
('人脸阈值 / 对齐复核阈值','0.5 / 0.1','观测筛选与复核提示'),
('文本词块 / CPU线程 / 随机种子','100 / 4 / 20260923','可复现参数')],[5.6,5.4,5.2],9.5)
p('预训练模型在准备阶段下载并登记修订号和逐文件哈希，正式提取阶段强制离线。源视频、标注表、代码、模型清单和配置均留有版本证据。断点续跑只有在处理指纹、源视频哈希和输出哈希均一致时跳过样本，避免直接沿用来源不明或已损坏的缓存。')
h('10.3 复现流程与交付组织',2)
p('依照requirements-lock.txt建立环境后，执行prepare_models.py准备模型，再运行run.ps1。入口依次调用算法检查、pipeline.py全量提取、summarize.py核验与汇总以及package_results.py打包。独立读取通过read_features.py完成，无需加载pickle对象。')
p('若仅调整聚合规则，可利用reaggregate.py从保存的原始帧特征重算，避免重新提取深度表示。本次结果记录了以float64累加器重算词级聚合的处理历史。压缩特征14.28 MiB不等于全部提交附件体积，最终提交还需统计代码、说明等文件，并按赛题统一限制检查总大小。')

# Page 12
page();h('11 模型评价与结论')
h('11.1 方法的合理性与适用范围',2)
p('本文把跨模态组织转化为“已知文本的时间定位”和“异步帧特征的区间统计”两个可计算子问题。前者由完整CTC拓扑约束的动态规划求解，后者由统一的时间重叠权重定义。两部分均具有明确输入、输出与失败规则，使特征生成不仅得到矩阵，还能提供矩阵位置与原始素材之间的对应记录。')
p('词级序列保留局部语言单位，均值与标准差同时描述区间内平均状态与波动。独立模态掩码能区分可用贡献和缺失位置，原始帧数据允许重新聚合。由于没有利用情感标签调整提取器，本问结果可作为基础表示与审计材料；但是否改善情感分类或回归，仍需后续模型与验证集实验支持。')
h('11.2 局限性与可实施的改进',2)
p('一是强制对齐对转写一致性敏感。18条样本平均分数低于0.1，应优先回听源音频，并对照保存的转写与模型直接识别文本区分原因；不能自行替换题目转写。二是10 Hz视觉采样和词级统计会压缩短暂动作及词内先后顺序，最大人脸规则也可能选错主体。可在后续人工复核后评估更密采样或说话人确认，但这些改进尚未在当前结果中完成。')
p('三是目前没有人工逐词边界真值，无法报告平均边界误差、对齐准确率或校准后的置信度。可预先制定抽样规则，标注起止时间，再计算误差及分数分层表现。四是当前768/50/110维特征与附件2的768/74/35维接口不同；问题二、三如采用附件2训练，应保持附件2规定的输入版本，不能将本问特征直接替换为同义输入。')
h('11.3 本问结论',2)
p('模型已为100条样本生成完整的词级三模态表示与追溯信息。1931个词均获得时间区间和有效语音贡献，1601个词具有有效视觉贡献；342个低分词及其他异常均保留复核标记。实现层面满足全量覆盖、可核验时序组织和复现记录要求，质量层面仍需人工边界标注与异常回看，不能将可用性统计外推为情感预测性能。')
h('参考资料',2)
refs=[
'[1] DistilBERT community. distilbert-base-uncased模型卡[EB/OL]. https://huggingface.co/distilbert/distilbert-base-uncased，访问日期2026-09-24.',
'[2] Meta. wav2vec2-base-960h模型卡[EB/OL]. https://huggingface.co/facebook/wav2vec2-base-960h，访问日期2026-09-24.',
'[3] audEERING. openSMILE Python Documentation[EB/OL]. https://audeering.github.io/opensmile-python/，访问日期2026-09-24.',
'[4] Google. Face landmark detection guide[EB/OL]. https://developers.google.com/edge/mediapipe/solutions/vision/face_landmarker，访问日期2026-09-24.',
'[5] 赛题材料. 复杂场景下多模态情感预测的数学建模与算法设计及附件1，2026.',
'[6] 问题一程序与运行记录. pipeline.py，alignment.py，summarize.py及results目录下全量特征、统计表和环境快照，2026.'
]
for text in refs:p(text,8.5,False).paragraph_format.line_spacing=Pt(12)

# Pages 13-15: all 100 sample records.
for group,(a,b) in enumerate([(0,34),(34,67),(67,100)]):
    page();h('12 全量样本特征提取结果' if group==0 else '12 全量样本特征提取结果续表')
    if group==0:
        p('表9按源表顺序列出全部100条样本。每条均包含文本T、语音A和视觉V，维度统一为768/50/110，粒度均为词级；因此将共有字段置于表注，逐行保留时长、词数及有效性。时长采用公共解码时间轴口径；A/V为语音和视觉有效词数，低分为q < 0.1的词数。',9.5)
    else:p('续表沿用相同字段与定义。全部记录的文本/语音/视觉维度均为768/50/110，对齐粒度均为词级；磁盘保存完整变长序列，无补齐和50词截断。',9.5)
    cap(f'表9  全部样本的逐条结果  第{a+1}至{b}条')
    tr=[]
    for idx,row in enumerate(ROWS[a:b],a+1):
        tr.append([idx,row['id'],f"{float(row['duration_s']):.3f}",row['word_count'],row['audio_valid_words'],row['vision_valid_words'],row['low_confidence_words'],'复核' if row['status']=='review_required' else '正常'])
    table(['序号','样本编号','时长/s','词数','A','V','低分','状态'],tr,[1,5.9,1.9,1.1,1.1,1.1,1.3,1.8],9)
    if group==2:
        p('合计：100条样本，787.5022秒，1931词，A = 1931，V = 1601，低分词342个，含复核提示样本69条。状态“正常”仅表示未触发当前配置的提示规则，不代表人工确认无误。全部词级时间区间、帧级数值及源索引保存于配套特征和元数据中。',9.5)
    elif group==1:p('注：复核状态还可能由数字读法、无人脸或多人脸提示触发，因此低分词数为0的样本仍可能需要复核。A/V掩码表示存在有效源帧贡献，不能据此认定词区间已被完整覆盖。',9)

doc.save(OUT)
print(OUT)
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'images',len(doc.inline_shapes),'equations',len(doc.element.xpath('.//m:oMath')))
