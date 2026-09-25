# -*- coding: utf-8 -*-
"""生成第一问 Word 论文：多模态情感特征提取与时序对齐的数学建模。"""
import csv
import json
import statistics
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent
RES = ROOT / 'results'
FIG = ROOT / 'figures'
OUT = Path(r'D:\Desktop\E\问题一_多模态特征提取与时序对齐建模.docx')

# ============ 数据载入 ============
def load(name):
    return json.loads((RES / name).read_text('utf-8'))

stats = load('validation.json')
env = load('run_environment.json')
example = load('example.json')
manifest = load('manifest.json')

with (RES / 'sample_summary.csv').open('r', encoding='utf-8-sig') as f:
    rows = list(csv.DictReader(f))
for r in rows:
    for k in ['duration_s', 'word_count', 'aligned_words', 'audio_valid_words',
              'vision_valid_words', 'alignment_confidence_mean', 'face_detection_rate']:
        r[k] = float(r[k])

durations = [r['duration_s'] for r in rows]
word_counts = [int(r['word_count']) for r in rows]
n_words = stats['words']
vision_cov = stats['vision_words'] / n_words
lowc_ratio = stats['low_confidence_words'] / n_words

# ============ 基础排版 helpers ============
SONG, HEI, TIMES = '宋体', '黑体', 'Times New Roman'

def set_run(run, cn=SONG, en=TIMES, size=12, bold=False, italic=False, color=None):
    run.font.name = en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)

def para(doc, text='', cn=SONG, en=TIMES, size=12, bold=False, italic=False,
         align=None, indent=True, spacing=22, before=0, after=0, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(spacing)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if align is not None:
        pf.alignment = align
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if text:
        run = p.add_run(text)
        set_run(run, cn, en, size, bold, italic, color)
    return p

def h1(doc, text):
    return para(doc, text, cn=HEI, en=TIMES, size=15, bold=True,
                align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, spacing=28, before=14, after=8)

def h2(doc, text):
    return para(doc, text, cn=HEI, en=TIMES, size=13, bold=True,
                align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, spacing=24, before=10, after=6)

def h3(doc, text):
    return para(doc, text, cn=HEI, en=TIMES, size=12, bold=True,
                align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, spacing=22, before=8, after=4)

def body(doc, text, **kw):
    kw.setdefault('spacing', 22)
    return para(doc, text, **kw)

def add_figure(doc, filename, caption, width_cm=15.0):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(FIG / filename), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    cap.paragraph_format.line_spacing = Pt(16)
    set_run(cap.add_run(caption), cn=SONG, size=10.5, bold=True)

# ============ OMML 公式 helpers ============
def mel(tag, parent=None):
    el = OxmlElement('m:' + tag)
    if parent is not None:
        parent.append(el)
    return el

def mr(parent, text, sty=None):
    r = mel('r', parent)
    if sty:
        rPr = mel('rPr', r)
        s = mel('sty', rPr)
        s.set(qn('m:val'), sty)
    t = mel('t', r)
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = text
    return r

def msub(parent, base_text, sub_text, base_sty=None, sub_sty=None):
    s = mel('sSub', parent)
    e = mel('e', s); mr(e, base_text, base_sty)
    sub = mel('sub', s); mr(sub, sub_text, sub_sty)
    return s

def msup(parent, base_text, sup_text, base_sty=None):
    s = mel('sSup', parent)
    e = mel('e', s); mr(e, base_text, base_sty)
    sup = mel('sup', s); mr(sup, sup_text)
    return s

def mfrac(parent, num_fn, den_fn):
    f = mel('f', parent)
    num = mel('num', f); num_fn(num)
    den = mel('den', f); den_fn(den)
    return f

def mdelim(parent, items_fn, beg='(', end=')'):
    d = mel('d', parent)
    dPr = mel('dPr', d)
    b = mel('begChr', dPr); b.set(qn('m:val'), beg)
    e2 = mel('endChr', dPr); e2.set(qn('m:val'), end)
    e = mel('e', d); items_fn(e)
    return d

def msum(parent, sub_fn, e_fn):
    nary = mel('nary', parent)
    naryPr = mel('naryPr', nary)
    c = mel('chr', naryPr); c.set(qn('m:val'), '∑')
    sh = mel('supHide', naryPr); sh.set(qn('m:val'), '1')
    sub = mel('sub', nary); sub_fn(sub)
    mel('sup', nary)
    e = mel('e', nary); e_fn(e)
    return nary

def msqrt(parent, e_fn):
    rad = mel('rad', parent)
    radPr = mel('radPr', rad)
    dh = mel('degHide', radPr); dh.set(qn('m:val'), '1')
    mel('deg', rad)
    e = mel('e', rad); e_fn(e)
    return rad

def msup_custom(parent):
    """创建 sSup 骨架，返回 (base_e, sup_e) 供自定义填充。"""
    s = mel('sSup', parent)
    base = mel('e', s)
    sup = mel('sup', s)
    return base, sup

def add_equation(doc, build_fn, number):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(24)
    pf.space_before = Pt(4)
    pf.space_after = Pt(4)
    pf.tab_stops.add_tab_stop(Cm(7.4), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(15.0), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t')
    om = OxmlElement('m:oMath')
    build_fn(om)
    p._p.append(om)
    p.add_run('\t(%d)' % number)
    return p

# ============ 表格 helpers ============
def set_table_borders(table, inner=False):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    spec = {'top': '12', 'bottom': '12', 'left': 'none', 'right': 'none',
            'insideV': 'none', 'insideH': '4' if inner else 'none'}
    for edge, val in spec.items():
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single' if val != 'none' else 'none')
        if val != 'none':
            el.set(qn('w:sz'), val)
            el.set(qn('w:color'), '808080' if edge == 'insideH' else '000000')
        borders.append(el)
    tblPr.append(borders)

def header_bottom_line(table):
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcB = OxmlElement('w:tcBorders')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:color'), '000000')
        tcB.append(bottom)
        tcPr.append(tcB)

def set_cell(cell, text, size=9, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, cn=SONG):
    cell.paragraphs[0].text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.alignment = align
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(size + 3)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    set_run(run, cn=cn, size=size, bold=bold)

def make_table(doc, data, widths=None, size=9, inner=True, align_first_left=False):
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for j, w in enumerate(widths):
            for i in range(len(data)):
                table.cell(i, j).width = Cm(w)
    for i, row in enumerate(data):
        for j, text in enumerate(row):
            al = WD_ALIGN_PARAGRAPH.LEFT if (align_first_left and j in (0, 1) and i > 0) else WD_ALIGN_PARAGRAPH.CENTER
            set_cell(table.cell(i, j), str(text), size=size,
                     bold=(i == 0), align=al)
    set_table_borders(table, inner=inner)
    header_bottom_line(table)
    return table

def add_table_caption(doc, caption):
    cap = doc.add_paragraph()
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(2)
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    cap.paragraph_format.line_spacing = Pt(16)
    set_run(cap.add_run(caption), cn=SONG, size=10.5, bold=True)

# ============ 页码 ============
def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run(run, size=9)
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)

# ============ 文档初始化 ============
doc = Document()
style = doc.styles['Normal']
style.font.name = TIMES
style.font.size = Pt(12)
style._element.rPr.rFonts.set(qn('w:eastAsia'), SONG)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.54)
sec.left_margin = sec.right_margin = Cm(2.8)
add_page_number(sec)

# ================= 标题与摘要 =================
para(doc, '复杂场景下多模态情感预测的数学建模与算法设计', cn=HEI, size=16, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, spacing=30, after=4)
para(doc, '——问题一：多模态情感特征提取与时序对齐的数学建模', cn=HEI, size=14, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, spacing=26, after=12)

para(doc, '摘　　要', cn=HEI, size=13, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, spacing=24, after=6)

abstract = ('针对复杂场景下多模态情感预测的第一问，本文建立了以"词"为时序锚点的三模态特征提取与跨模态时序'
            '对齐模型。模型分三层：第一层对文本、语音、视觉三模态分别提取词级语义特征（冻结 DistilBERT，'
            '子词均值池化，768 维）、帧级声学特征（PTS 感知 16 kHz 重采样与 openSMILE eGeMAPSv02 低层描述符，'
            '25 维）与帧级视觉特征（10 Hz 真实时间戳采样与 MediaPipe Face Landmarker 面部 blendshape 及头姿'
            '欧拉角，55 维）；第二层以 Wav2Vec2 CTC 声学后验为纽带，在完整 CTC 状态拓扑上用 Viterbi 动态规划'
            '将赛题转写强制对齐到音频时间轴，获得每个词约 20 ms 分辨率的时间区间及未校准置信度；第三层以词区间'
            '为窗口，对语音、视觉源帧按时间重叠长度加权聚合均值与标准差，形成与文本逐词对齐的三模态时序特征序列。'
            '模型对附件 1 全部 100 条样本完成无标签泄漏的全量提取：1931 个词 100% 获得词级时间区间与有效语音'
            '特征，视觉特征覆盖 82.9%，压缩特征共 14.3 MiB。全流程实现样本编号、模态文件与输出特征的一一对应，'
            '哈希溯源与结构自校验保证时序组织可核验，模型修订号锁定与指纹缓存保证方法可复现。')
body(doc, abstract, size=10.5, spacing=18)
para(doc, '', indent=False, spacing=8)
p = doc.add_paragraph()
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
p.paragraph_format.line_spacing = Pt(18)
p.paragraph_format.first_line_indent = Pt(21)
set_run(p.add_run('关键词：'), cn=HEI, size=10.5, bold=True)
set_run(p.add_run('多模态情感分析；特征提取；CTC 强制对齐；Viterbi 算法；时间重叠加权聚合；词级时序锚点'),
        size=10.5, bold=True)

doc.add_page_break()

# ================= 一、问题重述 =================
h1(doc, '一、问题重述')
h2(doc, '1.1 问题背景')
body(doc,
     '情感分析是具身智能与自然人机交互的关键支撑技术。真实交互中，人类情绪通过文本语义、语音韵律、面部表情与'
     '肢体动作等多通道信号协同表达。然而文本、语音、视觉模态在采样频率、特征维度与时间尺度上差异显著，特征提取'
     '与时序对齐的质量直接决定下游情感预测模型的性能上限。本题以 CMU-MOSEI 英文多模态情感数据集为基础，要求'
     '首先解决多模态特征提取与时序对齐的数学建模问题。')
h2(doc, '1.2 问题一的主要任务')
body(doc,
     '基于附件 1 提供的 100 条英文原始视频（解码时间轴总时长 %.1f 秒，单条 2.26～29.29 秒）及其转写文本与情感标注，'
     '建立文本、'
     '语音、视觉三模态情感特征提取与时序对齐的数学模型，明确原始样本的处理流程、各模态情感特征的定义与提取方法、'
     '跨模态时序对齐的规则与实现逻辑。自生成的特征文件须满足三项要求：一是原始样本覆盖完整性，样本编号、模态文件'
     '与输出特征一一对应；二是时序组织的可核验性，序列位置、有效长度、填充规则及其与原始素材的对应关系可追溯；'
     '三是方法的合理性与可复现性，特征提取工具、关键参数、处理日志和复现实验说明完备。论文还需给出全量结果汇总表'
     '与至少一个典型样本的时序对应验证。' % sum(durations))

# ================= 二、模型假设 =================
h1(doc, '二、模型假设与依据')
assumptions = [
    '（1）视频中的说话内容与赛题提供的转写文本在语义上基本一致，转写错误或数字读法差异属于局部扰动，可通过显式'
    '读法归一化与置信度标记处理，而不修改原始标注；',
    '（2）同一时刻画面中面积最大的人脸是当前表达主体（说话人）的近似，多人脸样本以复核标记保留证据，不强行认定；',
    '（3）情感相关的人脸表情与头部动作在 10 Hz 采样粒度下近似短时平稳，采样保持区间可以表征词级视觉状态；',
    '（4）特征提取工具（DistilBERT、Wav2Vec2、openSMILE、MediaPipe）参数全部冻结，不使用任何情感标签训练或'
    '微调特征提取器，保证问题一不发生标签泄漏；',
    '（5）音频时间轴上未观测到的间隙（容器缺口）不属于真实静音，特征统计时按缺失处理而非补零参与运算。',
]
for a in assumptions:
    body(doc, a, after=2)

# ================= 三、符号说明 =================
h1(doc, '三、符号说明')
add_table_caption(doc, '表 1　主要符号及含义')
symbols = [
    ['符号', '含义', '备注'],
    ['N', '样本总数', 'N = 100'],
    ['L', '单个样本的词数（valid_length）', '变长，5～65'],
    ['I_i = [s_i, e_i]', '第 i 个词的时间区间（秒）', 'CTC 强制对齐输出'],
    ['J_j = [g_j, f_j]', '第 j 个源特征帧的支持区间', '语音/视觉帧'],
    ['m_j ∈ {0,1}', '源帧有效性指示', '检测失败置 0'],
    ['w_ij', '词 i 与源帧 j 的时间重叠权重', '式（6）'],
    ['x_j', '源帧特征向量', '语音 25 维 / 视觉 55 维'],
    ['μ_i, σ_i', '词级加权均值与标准差', '式（7）（8）'],
    ['h_s', 'DistilBERT 子词最后隐层向量', '768 维'],
    ['S_i', '词 i 的子词下标集合', '分词器对齐'],
    ['ℓ', 'CTC 目标字符序列', '转写归一化后'],
    ['δ_t(s)', 'Viterbi 前向对数概率', '式（4）'],
    ['τ(k)', 'CTC 帧索引 k 到秒的时间映射', '式（3）'],
    ['Δt', 'CTC 帧步长', '约 0.02 s'],
    ['ρ_i', '词 i 的对齐置信度', '式（5），未校准'],
]
make_table(doc, symbols, widths=[4.2, 6.6, 4.2], size=10.5)
para(doc, '注：向量的上标 T 表示转置；I[·] 为指示函数。', size=9, indent=False, spacing=14, before=4)

# ================= 四、总体框架 =================
h1(doc, '四、总体框架与数据预处理')
h2(doc, '4.1 三层处理框架')
body(doc,
     '本文将问题一分解为"模态特征提取—词级强制对齐—时间重叠聚合"三个层次，如图 1 所示。核心思想是以"词"'
     '作为跨模态时序锚点：文本天然按词组织；语音与视觉的帧级特征通过强制对齐得到的词时间区间，被聚合到同一组'
     '词级位置上，从而形成文本 768 维、语音 50 维、视觉 110 维的三模态等长（同词数）特征序列。相比按固定网格'
     '（如 50 个位置）截断的方案，词级锚点不丢弃任何词语信息，天然支持变长序列，且每个位置的语义、声学与视觉'
     '证据严格对应同一发音事件，为问题二、三的融合与可解释性分析提供了可靠的时序基础。')
add_figure(doc, 'fig1_flow.png', '图 1　三模态特征提取与时序对齐总体流程', 14.5)

h2(doc, '4.2 样本登记与一致性校验')
body(doc,
     '以标签表 label-100.xlsx 为唯一登记入口，唯一编号定义为 id = video_id$_$clip_id。程序执行双向一致性校验：'
     '（i）标签表 100 行与 37 个子文件夹下的 100 个 mp4 文件集合完全对应，不多、不少、不重复；（ii）每行标签'
     '满足 label ∈ [−3, 3] 且 annotation 与 label 符号一致（0 仅属 Neutral）；（iii）源视频计算 SHA256 写入'
     '元数据。异常或极短样本一律保留原样进入处理流程，不做删除或替换，符合赛题"保留原始样本"的要求。')

# ================= 五、三模态特征提取 =================
h1(doc, '五、三模态特征提取模型')
h2(doc, '5.1 文本模态：词级上下文语义特征')
body(doc,
     '按正则表达式提取词元：保留撇号连接词（如 We\'re）、千分位数字（1,500）与序数词（10th），同时记录每个词'
     '在原始转写中的左闭右开字符偏移，保证"词—原文"双向可追溯。采用冻结参数的 DistilBERT（distilbert-base-'
     'uncased）作为文本编码器，以 100 词为一块、不截断地编码。设词 i 的子词下标集合为 S_i，h_s 为子词 s 在'
     '最后隐层的输出向量，则词级文本特征定义为子词均值：')
add_equation(doc, lambda om: (
    msub(om, 't', 'i'),
    mr(om, '='),
    mfrac(om,
          lambda num: mr(num, '1'),
          lambda den: (mr(den, '|'), msub(den, 'S', 'i'), mr(den, '|'))),
    msum(om, lambda sub: (mr(sub, 's'), mr(sub, '∈', sty='p'), msub(sub, 'S', 'i')),
         lambda e: msub(e, 'h', 's')),
), 1)
body(doc,
     '每个词得到 768 维上下文语义向量。分块内保留 chunk 首词索引与子词位置映射，编码超长时显式报错而非静默'
     '截断（当前数据最长 65 词，不会触发）。该设计使每个词的语义表征包含完整上下文，同时词级位置与对齐时间轴'
     '一一对应。')

h2(doc, '5.2 语音模态：帧级低层声学描述符')
body(doc,
     '用 PyAV 解码音频并重采样为 16 kHz 单声道。解码按 PTS（呈现时间戳）逐帧放置到以容器起点为 0 的公共时间轴上，'
     '记录每个采样点的"已观测"掩码——时间间隙被显式标记为未观测而非静音。在此基础上，用 openSMILE（eGeMAPSv02 '
     '配置，低层描述符级别）提取 25 维帧级声学特征，包括响度、谱斜率、MFCC（1～4）、基频（半音）、抖动、微颤、'
     '信噪比、前三个共振峰频率/带宽/相对幅度等，覆盖能量、韵律与音质三类情感声学线索。每帧由工具返回真实的起止'
     '时间戳；帧内若超过 5% 的样本点未观测，则该帧标记为无效。')

h2(doc, '5.3 视觉模态：面部动作与头部姿态特征')
body(doc,
     '按目标频率 10 Hz、以真实帧 PTS 为准对视频进行采样（不做均匀重采样，避免时间漂移），图像最长边缩放至 640 px。'
     '用 MediaPipe Face Landmarker 对每帧检测至多 3 张人脸（检测/存在置信度阈值 0.5），选择人脸框面积最大者，'
     '输出 52 维面部 blendshape 系数（眉毛、眼睛、面颊、下颌、口部等细粒度动作强度）与头部姿态：对面部变换矩阵'
     '的旋转部分做 SVD 正交化后提取 xyz 欧拉角，共 55 维原始视觉特征。检测失败帧特征置零且有效掩码置 0；多人脸'
     '帧保留人脸数量与所选框坐标供复核。每个采样帧的时间支持区间由相邻采样时刻的中点界定，并向左右各截断不超过'
     '0.075 s，防止跨越大的解码间隙。需要说明的是，blendshape 系数刻画的是动作强度而非情绪类别，也不是 OpenFace '
     '意义上的 FACS 动作单元（AU）。')

# ================= 六、时序对齐模型 =================
h1(doc, '六、跨模态时序对齐模型')
h2(doc, '6.1 CTC 强制对齐的数学描述')
body(doc,
     '跨模态对齐的目标是把"已知文本"映射到"音频时间轴"。设第 i 个样本的词序列为 w_1,…,w_L，对每个词做声学'
     '归一化（大写化；数字与序数词按显式映射表展开，如 2008→TWO THOUSAND EIGHT、10th→TENTH），词与词之间'
     '插入空白符 |，得到 CTC 目标字符序列：')
add_equation(doc, lambda om: (
    mr(om, 'ℓ'),
    mr(om, '='),
    mr(om, 'π', sty='p'),
    mdelim(om, lambda e: msub(e, 'w', '1')),
    mr(om, '|'),
    mr(om, 'π', sty='p'),
    mdelim(om, lambda e: msub(e, 'w', '2')),
    mr(om, '|⋯|'),
    mr(om, 'π', sty='p'),
    mdelim(om, lambda e: msub(e, 'w', 'L')),
), 2)
body(doc,
     '其中 π(w) 为词 w 的大写字符展开。冻结参数的 Wav2Vec2（facebook/wav2vec2-base-960h）输出每个 CTC 帧 t 上'
     '字符 c 的后验对数概率 log p_t(c)。强制对齐即在该后验矩阵上寻找一条从起始状态出发、恰好发射完整序列 ℓ 的'
     '最大概率路径。由于卷积前端感受野的限制，帧边界按卷积核中心定义：帧 k 的中心时刻与帧步长为')
add_equation(doc, lambda om: (
    mr(om, 'τ', sty='p'),
    mdelim(om, lambda e: mr(e, 'k')),
    mr(om, '='),
    msub(om, 'c', '0'),
    mr(om, '+'),
    mdelim(om, lambda e: (mr(e, 'k'), mr(e, '−'), mfrac(e, lambda n: mr(n, '1'), lambda d: mr(d, '2')))),
    mr(om, 'Δt'),
    mr(om, ',   '),
    mr(om, 'Δt', sty='p'),
    mr(om, '='),
    mfrac(om,
          lambda num: msub(num, 's', 'stride'),
          lambda den: msub(den, 'f', 's')),
), 3)
body(doc,
     '本模型中 Δt ≈ 0.02 s。需要强调，Δt 是模型的分辨率下界，不代表词边界误差小于 20 ms。')

h2(doc, '6.2 Viterbi 动态规划求解')
body(doc,
     '将 ℓ 展开为 blank 与字符交替的增广状态序列 (blank, ℓ_1, blank, ℓ_2, …)，允许三种转移：停留（原地重复'
     '发射）、前进 1 个状态、前进 2 个状态（仅当目标状态非 blank 且与前一字符不同，即跳过 blank）。记 δ_t(s) 为'
     '第 t 帧处于状态 s 的最大对数概率，则递推为：')
add_equation(doc, lambda om: (
    msub(om, 'δ', 't'),
    mdelim(om, lambda e: mr(e, 's')),
    mr(om, '='),
    mr(om, 'max', sty='p'),
    mdelim(om, lambda e: (
        msub(e, 'δ', 't−1'),
        mdelim(e, lambda e2: mr(e2, 's')),
        mr(e, ','),
        msub(e, 'δ', 't−1'),
        mdelim(e, lambda e2: (mr(e2, 's'), mr(e2, '−'), mr(e2, '1'))),
        mr(e, ','),
        msub(e, 'δ', 't−1'),
        mdelim(e, lambda e2: (mr(e2, 's'), mr(e2, '−'), mr(e2, '2'))),
        mr(e, '·'),
        mr(e, 'I', sty='p'),
        mdelim(e, lambda e2: (
            msub(e2, 'ℓ', 's'),
            mr(e2, '≠'),
            msub(e2, 'ℓ', 's−2'),
        ), beg='[', end=']'),
    ), beg='{', end='}'),
    mr(om, '+'),
    mr(om, 'log', sty='p'),
    msub(om, 'p', 't'),
    mdelim(om, lambda e: msub(e, 'ℓ', 's')),
), 4)
body(doc,
     '其中指示项 I[ℓ_s ≠ ℓ_(s−2)] 保证重复字符（如 ll、ss）之间必须经过 blank，这是完整 CTC 拓扑区别于简化'
     '贪心对齐的关键。终止时取最后两个状态中概率较大者回溯得到逐帧路径。若无可行路径（音频短于目标序列等），'
     '该样本的 alignment_mask 整体置 0 并记录错误，绝不伪造时间戳。')

h2(doc, '6.3 词级时间区间与置信度')
body(doc,
     '由回溯路径得到每个目标字符的发射帧集合，词 i 的时间区间取其首字符发射帧的左边界到末字符发射帧的右边界：'
     'I_i = [τ(u_i), τ(v_i+1)]，其中 u_i、v_i 为词 i 首末字符的发射帧索引。词 i 的对齐置信度定义为路径上该词'
     '各字符发射概率的几何平均：')
def _p_tk(parent):
    s = mel('sSub', parent)
    e = mel('e', s); mr(e, 'p')
    sub = mel('sub', s)
    msub(sub, 't', 'k')
add_equation(doc, lambda om: (
    msub(om, 'ρ', 'i'),
    mr(om, '='),
    mr(om, 'exp', sty='p'),
    mdelim(om, lambda eb: (
        mfrac(eb,
              lambda num: mr(num, '1'),
              lambda den: (mr(den, '|'), msub(den, 'C', 'i'), mr(den, '|'))),
        msum(eb,
             lambda sub: (mr(sub, 'k'), mr(sub, '∈', sty='p'), msub(sub, 'C', 'i')),
             lambda e2: (
                 mr(e2, 'log', sty='p'),
                 _p_tk(e2),
                 mdelim(e2, lambda e3: msub(e3, 'ℓ', 'k')),
             )),
    ), beg='[', end=']'),
), 5)
body(doc,
     '其中 C_i 为词 i 的字符发射帧集合（t_k 与 ℓ_k 为集合内第 k 帧的帧索引与对应目标字符）。置信度阈值 0.1 '
     '仅用作质量复核提示，未经边界真值校准，不作为精度指标；低置信度位置全部保留并单独打标，供人工回看。')

h2(doc, '6.4 逐词时间重叠加权聚合')
body(doc,
     '得到词区间后，语音与视觉的帧级特征按时间重叠长度加权聚合到词级。设词区间 I_i = [s_i, e_i]，源帧支持区间'
     'J_j = [g_j, f_j]，有效掩码 m_j，则重叠权重与词级特征为：')
add_equation(doc, lambda om: (
    msub(om, 'w', 'ij'),
    mr(om, '='),
    mr(om, 'max', sty='p'),
    mdelim(om, lambda e: (
        mr(e, '0'),
        mr(e, ','),
        mr(e, 'min', sty='p'),
        mdelim(e, lambda e2: (msub(e2, 'e', 'i'), mr(e2, ','), msub(e2, 'f', 'j'))),
        mr(e, '−'),
        mr(e, 'max', sty='p'),
        mdelim(e, lambda e2: (msub(e2, 's', 'i'), mr(e2, ','), msub(e2, 'g', 'j'))),
    )),
    mr(om, '·'),
    msub(om, 'm', 'j'),
), 6)
add_equation(doc, lambda om: (
    msub(om, 'μ', 'i'),
    mr(om, '='),
    mfrac(om,
          lambda num: msum(num, lambda sub: mr(sub, 'j'),
                           lambda e: (msub(e, 'w', 'ij'), msub(e, 'x', 'j'))),
          lambda den: msum(den, lambda sub: mr(sub, 'j'),
                           lambda e: msub(e, 'w', 'ij'))),
), 7)
def _sq_term(parent):
    """(x_j − μ_i)²"""
    base, sup = msup_custom(parent)
    mdelim(base, lambda eb: (msub(eb, 'x', 'j'), mr(eb, '−'), msub(eb, 'μ', 'i')))
    mr(sup, '2')
add_equation(doc, lambda om: (
    msub(om, 'σ', 'i'),
    mr(om, '='),
    msqrt(om, lambda e: mfrac(e,
        lambda num: msum(num, lambda sub: mr(sub, 'j'),
                         lambda e2: (msub(e2, 'w', 'ij'), _sq_term(e2))),
        lambda den: msum(den, lambda sub: mr(sub, 'j'),
                         lambda e2: msub(e2, 'w', 'ij')))),
), 8)
body(doc,
     '式（8）根号内为 Σ_j w_ij (x_j − μ_i)²。词级语音/视觉特征由均值与标准差拼接：a_i = [μ_i; σ_i] ∈ R^50、'
     'v_i = [μ_i; σ_i] ∈ R^110。聚合累加使用 float64，存储为 float32，避免大数值常量描述符因舍入产生伪标准差。'
     '当某词在某一模态下没有任何有效重叠帧时，该位置特征置 0、模态掩码置 0，即"显式缺失"而非"伪造观测"。'
     '词间未被任何词区间覆盖的时间段单独保存（unassigned_intervals），原始帧级特征与时间戳全部保留，词外信息'
     '并未从结果中丢弃。')

# ================= 七、特征文件规范 =================
h1(doc, '七、特征文件规范与存储设计')
h2(doc, '7.1 文件组织与字段定义')
body(doc,
     '每个样本输出一个 NPZ 特征文件（features/id.npz，禁用 pickle 以杜绝任意代码执行风险）与一个 JSON 元数据'
     '文件（metadata/id.json）。NPZ 主要字段见表 2。')
add_table_caption(doc, '表 2　NPZ 特征文件主要字段')
fields = [
    ['字段', '形状', '含义'],
    ['text', 'L × 768', '词级 DistilBERT 上下文特征'],
    ['audio', 'L × 50', '25 维声学描述符的词级均值+标准差'],
    ['vision', 'L × 110', '55 维面部特征的词级均值+标准差'],
    ['word_intervals', 'L × 2', '词级起止时间（秒）'],
    ['word_char_offsets', 'L × 2', '词在原始文本中的字符偏移'],
    ['alignment_confidence', 'L', '词级对齐置信度（未校准）'],
    ['text/audio/vision_mask', 'L', '三模态有效掩码'],
    ['alignment_mask / review_mask', 'L', '对齐有效 / 待复核标记'],
    ['audio_raw, audio_intervals', 'A × 25, A × 2', '原始声学帧及其时间'],
    ['vision_raw, vision_intervals', 'V × 55, V × 2', '原始视觉帧及其支持区间'],
    ['vision_pts, vision_frame_ids', 'V', '帧真实时间与原视频帧序号'],
    ['face_counts, face_boxes', 'V, V × 4', '检出人脸数与所选框坐标'],
    ['unassigned_intervals', 'G × 2', '词间未覆盖时间段'],
    ['valid_length / padding_mask', '标量 / L', '词数 / 填充标记（磁盘上全 0）'],
]
make_table(doc, fields, widths=[5.4, 3.2, 6.4], size=9.5, align_first_left=True)

h2(doc, '7.2 变长存储与填充规则')
body(doc,
     '磁盘保存变长完整序列：不截断到 50 词、不做零填充，valid_length 记录实际词数，padding_mask 在磁盘上恒为 '
     'False。批量读取时由 collate 函数按批内最大词数补齐并置 padding_mask=True，补齐位置三模态特征为 0 且模态'
     '掩码为 False，保证填充信息与真实观测严格可区分。该设计兼容任意词数（当前 5～65 词），也为问题二、三接入'
     '按附件 2 训练的模型提供了明确的接口转换规则。')

h2(doc, '7.3 数据溯源与完整性校验')
body(doc,
     '每条元数据记录：源视频 SHA256、特征文件 SHA256、处理指纹（配置+代码+模型清单的联合哈希）、逐词子词映射、'
     '聚合源帧索引（audio/vision_source_indices_per_word）、数字读法映射表与告警列表。汇总程序对 100 条样本执行'
     '结构校验（覆盖数、字符偏移回读、时间单调性、维度、有限性、掩码语义、哈希一致），并用独立实现的参考算法从'
     '原始帧特征重算全部词级聚合值进行数值复核，两套实现结果误差在 2×10⁻⁵ 以内。')

# ================= 八、结果分析 =================
h1(doc, '八、结果分析')
h2(doc, '8.1 全量提取结果概览')
body(doc,
     '100 条样本全部提取成功，结构校验零错误。样本时长分布与词数分布见图 2，情感极性分布见图 3：负向 18 条、'
     '中性 25 条、正向 57 条，与 label-100.xlsx 完全一致。全部样本共 1931 个词，词数最少的样本 5 个词、最多 '
     '65 个词，体现了说话长度的高度变异性。')
add_figure(doc, 'fig2_dist.png', '图 2　原始样本规模统计：(a) 时长分布；(b) 词数分布', 14.5)
add_figure(doc, 'fig6_label.png', '图 3　100 条样本情感极性分布', 9.0)

h2(doc, '8.2 三模态覆盖率分析')
body(doc,
     '图 4 给出三模态的有效覆盖率：词级时间区间（CTC 对齐）1931/1931 = 100.0%，有效语音特征 1931/1931 = 100.0%，'
     '有效视觉特征 1601/1931 = 82.9%。视觉未覆盖的 330 个词主要源于画外音、人脸遮挡、检测失败或人脸短暂离开画面，'
     '对应样本的人脸检出率分布见图 5(a)。这些位置以零特征+掩码 0 显式标记，不参与后续统计，为问题二的缺失模态'
     '鲁棒建模提供了真实分布的缺失样本。')
add_figure(doc, 'fig3_coverage.png', '图 4　三模态有效覆盖词比例', 11.0)

h2(doc, '8.3 提取质量与对齐置信度')
body(doc,
     '7889 个采样帧中 6462 帧检出有效人脸（81.9%），20 帧出现多人脸并已记录复核标记。对齐置信度方面，1931 个词'
     '中 342 个（' + ('%.1f' % (lowc_ratio * 100)) + '%）低于 0.1 的复核阈值；69 条样本含至少一个告警。按样本'
     '平均置信度统计，18 条样本低于 0.1，如图 5(b) 所示（多为深色点）：这些样本的赛题转写与声学模型直接识别文本'
     '差异较大，可能存在转写与音频不匹配或识别失败。本文的处理原则是：原转写与识别文本并列保存于 '
     'alignment_diagnostics.csv 供回听判断，不自行改写赛题文本；合法 CTC 路径仅说明算法找到受约束路径，不宣称'
     '该样本具有可靠的语义时间对应。')
add_figure(doc, 'fig5_quality.png', '图 5　提取质量分布：(a) 人脸检出率；(b) 对齐分数与低分词分布', 14.5)

h2(doc, '8.4 典型样本时序对应验证')
ex_dur = next(r['duration_s'] for r in rows if r['id'] == example['id'])
body(doc,
     '选取样本 ' + example['id'] + '（源视频 ' + example['source_video'] + '，时长 ' + ('%.2f' % ex_dur) +
     ' 秒，10 个词）作为典型样本：其全部词对齐成功、无低置信度词、人脸检出率 100% 且无多人脸帧。该样本转写为'
     '“' + example['raw_text'] + '”。图 6 展示了词级对齐、基频描述符、面部系数与采样帧在统一时间轴上的'
     '对应关系；表 3 列出逐词时间区间、置信度及聚合帧数。可以看到：实义词（think、social、justice、extremely、'
     'important）获得 24～40 帧语音证据与 3～5 帧视觉证据，置信度接近 1；句子末尾的 people 置信度 0.34，对应'
     '基频在 4.1 s 附近的突变，提示该词发音可能被截断或与转写不完全一致，已列入人工复核清单。')
add_figure(doc, 'fig4_alignment.png', '图 6　典型样本词—语音—视觉时序对齐（样本 %s）' % example['id'].replace('$_$', '$_$'), 15.0)

ex_words = example['word_rows']
tbl3 = [['词', '区间起点/s', '区间终点/s', '置信度', '语音帧数', '视觉帧数']]
for w in ex_words:
    tbl3.append([w['word'], '%.3f' % w['start_s'], '%.3f' % w['end_s'],
                 '%.3f' % w['score'], str(w['audio_frames']), str(w['visual_frames'])])
add_table_caption(doc, '表 3　典型样本逐词对齐结果（%s）' % example['id'])
make_table(doc, tbl3, widths=[2.6, 2.6, 2.6, 2.4, 2.4, 2.4], size=9.5)

h2(doc, '8.5 可复现性保障')
body(doc,
     '复现链条包括四层锁定：（i）模型层——prepare_models.py 记录三个预训练模型的仓库修订号与逐文件 SHA256，'
     '提取阶段强制离线；（ii）代码层——pipeline.py 与 alignment.py 的哈希写入运行环境快照；（iii）数据层——'
     '标签表与源视频的 SHA256 逐样本登记；（iv）处理指纹——配置、代码、模型清单的联合哈希写入每个样本元数据，'
     '--resume 模式仅在指纹完全匹配时跳过重算。聚合算法变更时可用 reaggregate.py 从保留的原始帧特征重算，无需'
     '重跑深度模型；本次交付已通过该脚本以 float64 累加器重算全部词级聚合，原特征、时间与掩码不变。')

# ================= 九、模型评价 =================
h1(doc, '九、模型评价与讨论')
h2(doc, '9.1 模型优点')
adv = [
    '（1）时序锚点统一：以词为跨模态锚点，三模态特征严格对齐同一发音事件，位置语义明确，可直接支撑问题三的'
    '"关键证据定位到词/帧"需求；',
    '（2）信息零丢弃：变长全序列存储，词外原始帧特征与未分配区间全部保留，避免固定网格截断带来的信息损失；',
    '（3）缺失显式化：三模态独立掩码 + 对齐复核掩码，把遮挡、无脸、对齐失败等质量信息编码为可计算的字段，'
    '直接服务于问题二的缺失鲁棒建模；',
    '（4）全程可核验：哈希溯源链、双实现数值复核、结构自校验与人工复核模板构成完整的质量证据链，符合赛题对'
    '"可核验性"与"可复现性"的三项硬性要求。',
]
for a in adv:
    body(doc, a, after=2)

h2(doc, '9.2 局限性与改进方向')
lims = [
    '（1）覆盖率不等于准确率：本文未做人工逐词边界标注，20 ms 帧步长只是分辨率下界；若需报告边界误差，应先对'
    '抽样片段人工听看标注（manual_review_template.csv 已预留字段）；',
    '（2）18 条样本平均对齐分数低于 0.1，疑似转写与音频不匹配，强制对齐可能落到错误位置，宜在论文中如实说明并'
    '结合回听结论区分对待；',
    '（3）视觉侧以"最大人脸"近似说话人，未做主动说话人确认；blendshape 亦非 FACS 动作单元，解释时需避免过度'
    '解读；10 Hz 采样会遗漏更短促的表情变化；',
    '（4）本文特征接口（词级 768/50/110 维）与附件 2 现成特征（50 位置、768/74/35 维）不同，问题二、三若使用'
    '附件 2 训练集，需通过附件 2 的统一接口训练，而本文特征作为问题一独立交付与对照验证。',
]
for a in lims:
    body(doc, a, after=2)

# ================= 参考文献 =================
h1(doc, '参考文献')
refs = [
    '[1] Sanh V, Debut L, Chaumond J, et al. DistilBERT, a distilled version of BERT: smaller, faster, cheaper '
    'and lighter[C]//NeurIPS EMC2 Workshop. 2019.',
    '[2] Baevski A, Zhou H, Mohamed A, et al. wav2vec 2.0: A framework for self-supervised learning of speech '
    'representations[C]//Advances in Neural Information Processing Systems. 2020, 33: 12449-12460.',
    '[3] Eyben F, Scherer K R, Schuller B W, et al. The Geneva Minimalistic Acoustic Parameter Set (eGeMAPS) for '
    'voice research and affective computing[J]. IEEE Transactions on Affective Computing, 2016, 7(2): 190-202.',
    '[4] Lugaresi C, Tang J, Nash H, et al. MediaPipe: A framework for building perception pipelines[J]. arXiv '
    'preprint arXiv:1906.08172, 2019.',
    '[5] Graves A, Fernández S, Gomez F, et al. Connectionist temporal classification: labelling unsegmented '
    'sequence data with recurrent neural networks[C]//Proceedings of the 23rd International Conference on Machine '
    'Learning. 2006: 369-376.',
    '[6] Zhang S, Yang Y, Chen C, et al. Deep learning-based multimodal emotion recognition from audio, visual, '
    'and text modalities: A systematic review[J]. Expert Systems with Applications, 2024, 237: 121692.',
    '[7] 王楠, 王淇, 欧阳丹彤. 基于知识蒸馏与动态调整机制的多模态情感分析模型[J]. 计算机学报, 2025, 48(8): '
    '1923-1942.',
]
for r in refs:
    para(doc, r, size=10.5, spacing=18, indent=False, after=2)

# ================= 附录 A =================
doc.add_page_break()
h1(doc, '附录 A　全量 100 条样本特征提取结果汇总表')
body(doc,
     '表 A1 按源表顺序列出全部 100 条样本的提取结果。文本/语音/视觉特征维度统一为 768/50/110 维，对齐粒度均为'
     '词级。"对齐词/语音有效/视觉有效"为该样本中获得词级时间区间、有效语音特征与有效视觉特征的词数。完整数值'
     '特征（features/*.npz）、逐词明细（word_alignment.csv）与原始帧级数据随附件提交。', size=10.5, spacing=18)
add_table_caption(doc, '表 A1　100 条样本特征提取结果汇总')

hdr = ['序号', '样本编号', '时长/s', '词数', '对齐词', '语音有效', '视觉有效', '平均置信度', '状态']
data = [hdr]
for idx, r in enumerate(rows, 1):
    status = '正常' if r['status'] == 'ok' else '待复核'
    data.append([str(idx), r['id'], '%.2f' % r['duration_s'], str(int(r['word_count'])),
                 str(int(r['aligned_words'])), str(int(r['audio_valid_words'])),
                 str(int(r['vision_valid_words'])), '%.3f' % r['alignment_confidence_mean'], status])
tbl = make_table(doc, data, widths=[1.0, 3.6, 1.5, 1.2, 1.4, 1.6, 1.6, 1.7, 1.4], size=8)
# 統計行
mean_dur = statistics.mean(durations)
mean_words = statistics.mean(word_counts)
mean_conf = statistics.mean([r['alignment_confidence_mean'] for r in rows])
sum_aligned = sum(int(r['aligned_words']) for r in rows)
sum_audio = sum(int(r['audio_valid_words']) for r in rows)
sum_vision = sum(int(r['vision_valid_words']) for r in rows)
last = data[-1]
row_cells = tbl.add_row().cells
vals = ['均值/合计', '%d 条' % len(rows), '%.2f' % mean_dur, '%.1f' % mean_words,
        str(sum_aligned), str(sum_audio), str(sum_vision), '%.3f' % mean_conf, '—']
for c, v in zip(row_cells, vals):
    set_cell(c, v, size=8, bold=True)
para(doc, '注："待复核"表示样本含低置信度对齐词、无人脸词区间或多人脸帧等告警，证据均已保留在元数据中；'
          '状态与告警明细见 sample_summary.csv。', size=9, indent=False, spacing=14, before=4)

# ================= 附录 B =================
h1(doc, '附录 B　复现运行流程与版本清单')
h2(doc, 'B.1 运行环境与模型版本')
env_tbl = [['组件', '版本/修订号'],
           ['Python', env['python']],
           ['numpy / scipy / pandas', '%s / %s / %s' % (env['packages']['numpy'], env['packages']['scipy'], env['packages']['pandas'])],
           ['PyAV（视频/音频解码）', env['packages']['av']],
           ['openSMILE（eGeMAPSv02）', env['packages']['opensmile']],
           ['MediaPipe Face Landmarker', env['packages']['mediapipe']],
           ['transformers / torch', '%s / %s' % (env['packages']['transformers'], env['packages']['torch'])],
           ['DistilBERT 模型修订号', env['models']['models']['text']['revision'][:12]],
           ['Wav2Vec2 模型修订号', env['models']['models']['alignment']['revision'][:12]],
           ['MediaPipe 人脸关键点模型', 'face_landmarker v1（float16）'],
           ['关键参数', '16 kHz；10 Hz；640 px；检测阈值 0.5；复核阈值 0.1；seed 20260923']]
make_table(doc, env_tbl, widths=[6.6, 8.4], size=9.5, align_first_left=True)

h2(doc, 'B.2 一键复现命令（Windows PowerShell）')
cmds = [
    'python -m venv .venv',
    '.venv/Scripts/python.exe -m pip install -r question1/requirements-lock.txt',
    '.venv/Scripts/python.exe question1/prepare_models.py        # 下载并哈希校验预训练模型',
    '.venv/Scripts/python.exe question1/test_alignment.py        # 算法单元测试',
    '.venv/Scripts/python.exe question1/pipeline.py --data-root E题 --resume   # 全量提取',
    '.venv/Scripts/python.exe question1/summarize.py --data-root E题           # 校验与报告',
    '.venv/Scripts/python.exe question1/package_results.py       # 打包（≤50 MB 校验）',
]
for c in cmds:
    para(doc, c, size=9.5, spacing=15, indent=False, after=1, cn=SONG)
body(doc,
     'CPU 可运行，无需 GPU。--limit 2 可先小规模验证；--resume 依据处理指纹（配置+代码+模型清单哈希）与源文件'
     '哈希决定是否跳过，配置或代码变更将触发重算。全部 100 条样本在 4 线程 CPU 上的提取耗时约 15 分钟。', size=10.5)

doc.save(str(OUT))
print('saved:', OUT)
